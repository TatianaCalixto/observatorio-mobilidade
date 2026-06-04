"""
Gera a planilha XLSX de planejamento das sprints do projeto Observatório de Mobilidade.

A planilha é a "documentação verdade" que o Claude Code usará para executar as tarefas
em sequência, marcar progresso e registrar impedimentos.

Abas:
  1. Instruções        -> Como o Claude Code deve usar a planilha
  2. Visão Geral       -> Status macro de cada sprint
  3. Backlog           -> Backlog detalhado (uma linha por tarefa)
  4. Plano de Testes   -> Checklist de testes por feature
  5. Impedimentos      -> Log de bloqueios / dúvidas para o humano
  6. Decisões          -> Log de decisões técnicas tomadas
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ============================================================
# Estilos reutilizáveis
# ============================================================

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E78")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1F4E78")
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SPRINT_FILL = PatternFill("solid", fgColor="D9E1F2")

STATUS_OPTIONS = [
    "Pendente",
    "Em andamento",
    "Bloqueada",
    "Em teste",
    "Concluída",
    "Cancelada",
]

PRIORIDADE_OPTIONS = ["Alta", "Média", "Baixa"]
TIPO_OPTIONS = ["Backend", "Frontend", "Mobile", "Infra", "Banco", "Teste", "Documentação", "Deploy"]


# ============================================================
# Conteúdo: Sprints e Tarefas
# ============================================================

SPRINTS = [
    {
        "numero": 1,
        "nome": "Sprint 1 — Fundação do Repositório",
        "objetivo": "Preparar repositório, ambiente Python reprodutível (uv), lint, "
                    "pre-commit, testes, Makefile e CI mínimo. Sem essa base nada do "
                    "resto roda de forma reprodutível.",
        "entregaveis": "Repo versionado, ambiente travado por lockfile, ruff + pre-commit, "
                       "pytest configurado, Makefile com comandos padrão, GitHub Actions "
                       "rodando lint + testes em cada push.",
    },
    {
        "numero": 2,
        "nome": "Sprint 2 — Ingestão (Extract & Load)",
        "objetivo": "Coletar GTFS (SPTrans), clima (INMET), população (IBGE/SIDRA) e dados "
                    "da Base dos Dados, gravando em camada RAW (Parquet particionado) de "
                    "forma idempotente, e carregar no DuckDB.",
        "entregaveis": "Scripts de ingestão idempotentes, camada RAW em Parquet, dados "
                       "carregados no DuckDB, validação de schema na entrada, logging.",
    },
    {
        "numero": 3,
        "nome": "Sprint 3 — Modelagem dbt (Transform)",
        "objetivo": "Projeto dbt-duckdb com camadas staging (limpeza/tipagem), intermediate "
                    "(joins clima/população) e marts (modelo dimensional: fato de viagens/"
                    "atrasos + dimensões linha, parada, tempo, clima), com testes de qualidade.",
        "entregaveis": "dbt build verde, modelos nas 3 camadas, testes not_null/unique/"
                       "relationships passando, documentação dbt gerada.",
    },
    {
        "numero": 4,
        "nome": "Sprint 4 — Análise Exploratória & Métricas",
        "objetivo": "Responder a pergunta de negócio via SQL/notebook: gargalos da rede, "
                    "sazonalidade e correlação clima × atraso. Consolidar marts analíticos "
                    "prontos para o dashboard.",
        "entregaveis": "Marts analíticos, relatório com insights, métrica de atraso/demanda "
                       "definida e testada (regra de negócio crítica).",
    },
    {
        "numero": 5,
        "nome": "Sprint 5 — Machine Learning",
        "objetivo": "Feature engineering a partir dos marts, baseline honesto, modelo "
                    "principal (XGBoost) com split temporal, tracking de experimentos "
                    "(MLflow) e serialização do modelo para o app.",
        "entregaveis": "Pipeline de features, baseline + modelo avaliados com split temporal, "
                       "experimentos registrados no MLflow, modelo serializado.",
    },
    {
        "numero": 6,
        "nome": "Sprint 6 — App Streamlit",
        "objetivo": "Dashboard com KPIs, mapa interativo da rede (gargalos), explorador de "
                    "linhas e seção de previsões consumindo o modelo treinado.",
        "entregaveis": "App Streamlit local funcional consumindo marts (DuckDB) e o modelo, "
                       "com estados vazio/carregando/erro tratados.",
    },
    {
        "numero": 7,
        "nome": "Sprint 7 — Orquestração & CI/CD (Hardening)",
        "objetivo": "Pipeline Prefect orquestrando ingestão → dbt → treino, com schedule e "
                    "retries. CI rodando lint, pytest e dbt test. Dockerfile para execução "
                    "portátil. Cobertura mínima de testes.",
        "entregaveis": "Pipeline Prefect agendado, GitHub Actions verde (lint + pytest + "
                       "dbt test), Dockerfile, cobertura ≥ 80% no código Python.",
    },
    {
        "numero": 8,
        "nome": "Sprint 8 — Documentação & Publicação",
        "objetivo": "README com storytelling da pergunta de negócio, diagrama de arquitetura, "
                    "GIF do app, instruções de reprodução, ADRs leves e publicação do app.",
        "entregaveis": "App público (Streamlit Community Cloud), README de portfólio, ADRs, "
                       "screenshots/GIF, projeto reproduzível por terceiros via make.",
    },
]


# Cada item: (sprint, id_tarefa, tipo, titulo, descricao, criterios_aceitacao,
#             dependencias, testes_obrigatorios, prioridade)

TAREFAS = [
    # ---------------- Sprint 1 — Fundação ----------------
    (1, "S01-T01", "Infra",
     "Criar repositório e estrutura de pastas",
     "Inicializar repositório Git. Criar estrutura: ingestion/, transform/ (dbt), "
     "ml/, app/, analysis/, tests/, docs/, data/{raw,staging}/. Adicionar .gitignore "
     "para Python, dados locais (data/, *.duckdb, *.parquet), mlruns/, .venv e IDEs.",
     "Repositório criado; .gitignore cobre dados locais, .venv, mlruns/, .duckdb, "
     "__pycache__, .idea/, .vscode/; README raiz presente descrevendo o projeto.",
     "—",
     "Validação manual: clonar em outra pasta e confirmar que dados e segredos não entram.",
     "Alta"),
    (1, "S01-T02", "Infra",
     "Ambiente Python reprodutível com uv",
     "Configurar pyproject.toml com uv. Dependências base: duckdb, pandas, polars, "
     "pyarrow, requests, httpx, python-dotenv, pytest, ruff. Gerar lockfile (uv.lock).",
     "uv sync instala sem erros; lockfile versionado; Python 3.12 fixado.",
     "S01-T01",
     "uv sync em pasta limpa sem erros; uv run python -c 'import duckdb, pandas' OK.",
     "Alta"),
    (1, "S01-T03", "Infra",
     "Lint e pre-commit",
     "Configurar ruff (lint + format) no pyproject. Adicionar .pre-commit-config.yaml "
     "com hooks de ruff, ruff-format e trailing-whitespace.",
     "ruff check . sem violações no esqueleto; pre-commit install funciona e bloqueia "
     "commit com violação.",
     "S01-T02",
     "ruff check . limpo; rodar pre-commit run --all-files e validar gates.",
     "Média"),
    (1, "S01-T04", "Teste",
     "Configurar pytest e coverage",
     "Criar tests/ com conftest.py e pytest.ini/pyproject [tool.pytest]. Configurar "
     "pytest-cov. Adicionar teste smoke validando import dos pacotes do projeto.",
     "pytest roda e exibe 1 teste passando; pytest --cov gera relatório.",
     "S01-T02",
     "Teste smoke (import dos módulos); pytest --cov mostra relatório de cobertura.",
     "Alta"),
    (1, "S01-T05", "Infra",
     "Makefile com comandos padrão",
     "Criar Makefile com alvos: setup (uv sync), lint, test, ingest, dbt-build, train, "
     "app, all. Documentar cada alvo.",
     "make setup, make lint e make test funcionam de ponta a ponta na máquina limpa.",
     "S01-T03, S01-T04",
     "Executar make setup && make lint && make test em pasta limpa.",
     "Média"),
    (1, "S01-T06", "Infra",
     "CI mínimo (GitHub Actions)",
     "Criar .github/workflows/ci.yml que roda uv sync, ruff check e pytest em cada "
     "push/PR.",
     "Workflow verde em push de branch nova; falha vermelha quando teste é quebrado "
     "de propósito.",
     "S01-T05",
     "Push quebrando assert para verificar falha do CI; depois corrigir.",
     "Média"),

    # ---------------- Sprint 2 — Ingestão ----------------
    (2, "S02-T01", "Backend",
     "Config central de caminhos e settings",
     "Criar módulo de config (paths das camadas raw/staging, URLs das fontes, janela de "
     "datas) lido de .env via python-dotenv. Fixar o recorte temporal do projeto: "
     "DATA_INICIO=2025-06-01 e DATA_FIM=2026-05-31 (últimos 12 meses). Criar .env.example.",
     "Config carrega valores do .env incluindo a janela de 12 meses; ausência de variável "
     "obrigatória gera erro claro; .env.example documenta todas as variáveis.",
     "S01-T02",
     "Unit test: carregar config com .env de teste e validar campos e erro em ausência.",
     "Alta"),
    (2, "S02-T02", "Backend",
     "Ingestão GTFS (SPTrans) → RAW Parquet",
     "Baixar/ler o feed GTFS estático de São Paulo (stops, routes, trips, stop_times) e "
     "gravar cada tabela em data/raw/gtfs/ como Parquet particionado. Idempotente.",
     "Arquivos GTFS principais salvos em Parquet; reexecução não duplica dados; log do "
     "volume ingerido por tabela.",
     "S02-T01",
     "Teste com fixture GTFS reduzida: parsing correto, schema esperado, idempotência "
     "(rodar 2x gera o mesmo resultado).",
     "Alta"),
    (2, "S02-T03", "Backend",
     "Ingestão INMET (clima) → RAW Parquet",
     "Coletar dados climáticos por estação (precipitação, temperatura) para a região de "
     "SP no recorte de 12 meses (2025-06-01 a 2026-05-31) e gravar em data/raw/clima/ "
     "como Parquet. Idempotente.",
     "Dados de clima salvos em Parquet com colunas tipadas (data, estação, métricas), "
     "limitados à janela de 12 meses; reexecução idempotente.",
     "S02-T01",
     "Teste com fixture de CSV INMET: parsing, tipagem de data/numéricos, idempotência.",
     "Alta"),
    (2, "S02-T04", "Backend",
     "Ingestão IBGE/SIDRA (população) → RAW Parquet",
     "Consumir API SIDRA para população por distrito/região de SP e gravar em "
     "data/raw/populacao/ como Parquet.",
     "Tabela de população salva com chave de região compatível com o GTFS/marts.",
     "S02-T01",
     "Teste com resposta SIDRA mockada: parsing do JSON e schema de saída.",
     "Média"),
    (2, "S02-T05", "Banco",
     "Carga das camadas RAW no DuckDB (idempotente)",
     "Criar script que registra/recarrega os Parquet de data/raw/ como tabelas/views no "
     "DuckDB (mobilidade.duckdb). Recriação idempotente (CREATE OR REPLACE).",
     "DuckDB contém tabelas raw_gtfs_*, raw_clima, raw_populacao; reexecução não "
     "duplica; SELECT count(*) coerente com os Parquet.",
     "S02-T02, S02-T03, S02-T04",
     "Teste de integração: carregar fixtures no DuckDB temporário e validar contagens e "
     "idempotência (carregar 2x).",
     "Alta"),
    (2, "S02-T06", "Teste",
     "Suite de validação de ingestão",
     "Consolidar validações de schema (colunas obrigatórias, tipos) e idempotência das "
     "ingestões em uma suíte reutilizável.",
     "Falha clara quando o schema de origem muda; idempotência coberta para cada fonte.",
     "S02-T05",
     "Testes de schema por fonte + teste de idempotência consolidado.",
     "Média"),

    # ---------------- Sprint 3 — Modelagem dbt ----------------
    (3, "S03-T01", "Infra",
     "Setup dbt-duckdb",
     "Inicializar projeto dbt em transform/, configurar profiles.yml apontando para o "
     "DuckDB do projeto, definir camadas (staging, intermediate, marts) em dbt_project.yml.",
     "dbt debug verde; dbt run vazio executa sem erro contra o DuckDB.",
     "S02-T05",
     "dbt debug e dbt parse sem erros no CI.",
     "Alta"),
    (3, "S03-T02", "Backend",
     "Modelos staging (limpeza e tipagem)",
     "Criar models/staging/ com um stg_* por fonte (gtfs stops/routes/trips/stop_times, "
     "clima, população): renomear colunas, tipar, padronizar chaves.",
     "Modelos stg_* materializados como views; nomes/tipos padronizados; sem nulos em "
     "chaves declaradas.",
     "S03-T01",
     "dbt test com not_null/unique nas chaves dos stg_*; dbt build verde.",
     "Alta"),
    (3, "S03-T03", "Backend",
     "Modelos intermediate (joins e enriquecimento)",
     "Criar models/intermediate/ unindo horários planejados (GTFS) com clima por data e "
     "população por região, preparando a base de viagens/atrasos restrita à janela de 12 "
     "meses do projeto (2025-06-01 a 2026-05-31).",
     "Modelos int_* unem as fontes sem explosão de linhas; relationships válidas entre "
     "as chaves.",
     "S03-T02",
     "dbt test relationships entre int_* e stg_*; verificação de contagem (sem fan-out).",
     "Alta"),
    (3, "S03-T04", "Backend",
     "Marts dimensionais (fato + dimensões)",
     "Criar models/marts/: fct_viagens (com métrica de atraso/headway), dim_linha, "
     "dim_parada, dim_tempo, dim_clima. Documentar colunas no schema.yml.",
     "Marts materializados como tabelas; grão do fato documentado; dimensões com chave "
     "única; descrições no schema.yml.",
     "S03-T03",
     "dbt test unique/not_null nas dimensões; testes de relationships do fato com as "
     "dimensões; dbt docs generate.",
     "Alta"),
    (3, "S03-T05", "Teste",
     "Qualidade dbt e build verde end-to-end",
     "Garantir cobertura de testes dbt em todas as camadas e que dbt build roda limpo do "
     "zero. Integrar dbt build no CI.",
     "dbt build (run + test) verde do zero; CI executa dbt build em container.",
     "S03-T04",
     "dbt build completo no CI; falha proposital em um teste valida o gate.",
     "Alta"),

    # ---------------- Sprint 4 — Análise ----------------
    (4, "S04-T01", "Backend",
     "Definição e cálculo da métrica de atraso/demanda (regra crítica)",
     "Definir formalmente a métrica central (ex.: atraso = real - planejado por trecho, "
     "ou densidade de viagens por linha/hora) em um mart analítico, com documentação do "
     "cálculo.",
     "Métrica calculada de forma determinística e documentada; valores conferem em casos "
     "manuais conhecidos.",
     "S03-T04",
     "TESTES OBRIGATÓRIOS: casos manuais com resultado esperado; sinal correto; "
     "tratamento de dados faltantes; teste que trava regressão no cálculo da métrica.",
     "Alta"),
    (4, "S04-T02", "Backend",
     "Análise de gargalos e sazonalidade",
     "Queries/notebook identificando linhas/paradas com maiores atrasos e padrões "
     "sazonais (dia da semana, hora, mês).",
     "Top gargalos e padrões sazonais reproduzíveis a partir dos marts; resultados "
     "salvos como mart analítico ou notebook versionado.",
     "S04-T01",
     "Teste de integração validando que as queries retornam o schema esperado e não "
     "quebram com período vazio.",
     "Média"),
    (4, "S04-T03", "Backend",
     "Correlação clima × atraso",
     "Analisar relação entre variáveis climáticas (chuva/temperatura) e atrasos, "
     "consolidando em tabela analítica.",
     "Tabela de correlação/agregação por condição climática gerada; método documentado.",
     "S04-T02",
     "Teste do agregado climático com fixture (sem chuva vs com chuva).",
     "Média"),
    (4, "S04-T04", "Documentação",
     "Relatório de insights (pergunta de negócio)",
     "Escrever relatório (notebook ou MD) respondendo a pergunta de negócio com 3+ "
     "insights e gráficos.",
     "Relatório responde explicitamente a pergunta do blueprint com evidências dos marts.",
     "S04-T03",
     "Revisão: relatório reproduzível a partir dos marts (re-rodar gera os mesmos números).",
     "Média"),

    # ---------------- Sprint 5 — Machine Learning ----------------
    (5, "S05-T01", "Backend",
     "Feature engineering a partir dos marts",
     "Construir dataset de features (linha, parada, dia da semana, hora, clima, "
     "população) e o alvo (atraso/demanda) a partir dos marts, com split temporal "
     "definido (treino/validação por data).",
     "Dataset reprodutível e versionável; sem vazamento temporal (features não usam "
     "futuro); split por data documentado.",
     "S04-T01",
     "TESTES OBRIGATÓRIOS: shape e colunas do dataset; ausência de vazamento temporal "
     "(corte de data respeitado); determinismo com seed fixa.",
     "Alta"),
    (5, "S05-T02", "Backend",
     "Baseline honesto",
     "Implementar baseline simples (média por linha/hora ou regressão linear) e calcular "
     "métricas (MAE/RMSE) no conjunto de validação temporal.",
     "Baseline treina e reporta métricas reprodutíveis; serve de referência para o "
     "modelo principal.",
     "S05-T01",
     "Teste: métricas do baseline dentro de faixa esperada em fixture; reprodutível "
     "com seed.",
     "Alta"),
    (5, "S05-T03", "Infra",
     "Tracking de experimentos com MLflow",
     "Configurar MLflow (tracking local em mlruns/) para registrar params, métricas e "
     "artefatos de cada execução de treino.",
     "Cada treino gera um run no MLflow com params e métricas; runs comparáveis na UI.",
     "S05-T02",
     "Teste: execução de treino registra run com as métricas esperadas (tracking URI "
     "temporário).",
     "Média"),
    (5, "S05-T04", "Backend",
     "Modelo principal (XGBoost) e avaliação",
     "Treinar XGBoost com o mesmo split temporal, comparar com baseline, registrar no "
     "MLflow e selecionar o melhor.",
     "Modelo principal supera (ou iguala com justificativa) o baseline; avaliação "
     "honesta documentada; melhor run identificado.",
     "S05-T03",
     "Teste: pipeline de treino roda em fixture e produz métricas; comparação "
     "baseline×modelo registrada.",
     "Alta"),
    (5, "S05-T05", "Backend",
     "Serialização do modelo para o app",
     "Persistir o modelo selecionado (e o pré-processamento) em artefato carregável pelo "
     "app, com metadados (versão, métricas, data).",
     "Artefato salvo e recarregável; predição em amostra reproduz a métrica registrada.",
     "S05-T04",
     "Teste: carregar artefato e prever amostra conhecida com resultado estável.",
     "Alta"),

    # ---------------- Sprint 6 — App Streamlit ----------------
    (6, "S06-T01", "Frontend",
     "Estrutura do app Streamlit + conexão DuckDB",
     "Criar app/ com Streamlit; camada de dados que lê os marts do DuckDB (read-only) "
     "com cache; layout base e navegação por páginas.",
     "App sobe localmente; lê marts sem travar; estados de carregamento tratados.",
     "S03-T04",
     "Smoke test do app (streamlit não quebra no import; funções de dados retornam "
     "DataFrame).",
     "Alta"),
    (6, "S06-T02", "Frontend",
     "Dashboard de KPIs e análises",
     "Página com KPIs (atraso médio, top linhas), filtros por período/linha e gráficos "
     "consumindo os marts analíticos.",
     "KPIs e gráficos refletem os marts; filtros funcionam; estado vazio tratado.",
     "S06-T01, S04-T02",
     "Teste das funções de agregação que alimentam os gráficos (entrada→saída).",
     "Média"),
    (6, "S06-T03", "Frontend",
     "Mapa interativo de gargalos",
     "Mapa (pydeck/folium) plotando paradas/linhas com intensidade de atraso a partir "
     "das coordenadas do GTFS.",
     "Mapa renderiza pontos georreferenciados; intensidade reflete a métrica; sem erro "
     "com dados faltantes de coordenada.",
     "S06-T02",
     "Teste da função que prepara o GeoDataFrame/camada do mapa.",
     "Média"),
    (6, "S06-T04", "Frontend",
     "Seção de previsões (consumo do modelo)",
     "Página que carrega o modelo serializado e gera previsão para entradas selecionadas "
     "(linha, dia, hora, clima).",
     "Previsão exibida para entrada do usuário; carregamento do modelo cacheado; erro "
     "tratado se artefato ausente.",
     "S05-T05, S06-T01",
     "Teste: função de predição do app retorna valor para entrada válida usando "
     "artefato de fixture.",
     "Alta"),
    (6, "S06-T05", "Teste",
     "Smoke tests do app consolidados",
     "Garantir que cada página importa e que as funções de dados/predição têm cobertura "
     "mínima.",
     "Suíte de smoke do app verde no CI.",
     "S06-T04",
     "Smoke tests por página + cobertura das funções de dados do app.",
     "Média"),

    # ---------------- Sprint 7 — Orquestração & CI/CD ----------------
    (7, "S07-T01", "Infra",
     "Pipeline Prefect (ingestão → dbt → treino)",
     "Criar flow Prefect orquestrando: ingestão (S2) → carga DuckDB → dbt build (S3) → "
     "treino/registro (S5). Tasks com retries e logging.",
     "flow roda end-to-end localmente reconstruindo marts e modelo; retries configurados "
     "em tasks de rede.",
     "S05-T04, S03-T05",
     "Teste do flow em modo reduzido (fixtures) validando a ordem e a conclusão das tasks.",
     "Alta"),
    (7, "S07-T02", "Infra",
     "Schedule do pipeline",
     "Configurar deployment/schedule do flow (ex.: diário) documentando como rodar local "
     "ou agendado.",
     "Schedule definido e documentado; execução manual do deployment funciona.",
     "S07-T01",
     "Validação: disparar o deployment manualmente e confirmar execução completa.",
     "Baixa"),
    (7, "S07-T03", "Infra",
     "CI completo (lint + pytest + dbt test)",
     "Estender GitHub Actions para rodar ruff, pytest com cobertura e dbt build em "
     "container, falhando abaixo do alvo de cobertura.",
     "CI verde em push; falha quando cobertura < 80% ou dbt test quebra.",
     "S07-T01",
     "Quebra proposital de teste/dbt valida os gates do CI.",
     "Alta"),
    (7, "S07-T04", "Deploy",
     "Dockerfile e execução portátil",
     "Criar Dockerfile (multi-stage) que instala o ambiente via uv e roda o pipeline/"
     "app; documentar comandos.",
     "docker build conclui; container roda o flow e/ou o app sem erro.",
     "S07-T03",
     "Build local + run do container executando make all em modo reduzido.",
     "Média"),
    (7, "S07-T05", "Teste",
     "Cobertura mínima e regressão end-to-end",
     "Elevar cobertura do código Python para ≥ 80% e garantir que a regressão da métrica "
     "(S04-T01) e do dataset (S05-T01) seguem verdes.",
     "Cobertura ≥ 80% global; baterias de regressão de métrica e de features verdes.",
     "S07-T03",
     "Relatório de cobertura no CI; reexecução das regressões críticas.",
     "Alta"),

    # ---------------- Sprint 8 — Documentação & Publicação ----------------
    (8, "S08-T01", "Documentação",
     "README de portfólio com storytelling",
     "Escrever README contando a pergunta de negócio, arquitetura (diagrama), stack, "
     "resultados/insights e como reproduzir (make).",
     "README permite a um terceiro entender e reproduzir o projeto; diagrama presente.",
     "S07-T05",
     "Revisão: seguir o README do zero em pasta nova reproduz pipeline + app.",
     "Alta"),
    (8, "S08-T02", "Deploy",
     "Publicar o app (Streamlit Community Cloud)",
     "Publicar o app com dados/artefatos necessários; ajustar caminhos e segredos para o "
     "ambiente público.",
     "App acessível por link público respondendo à pergunta de negócio.",
     "S08-T01",
     "Smoke manual no app publicado (carrega dashboard, mapa e previsão).",
     "Alta"),
    (8, "S08-T03", "Documentação",
     "ADRs leves e log de decisões",
     "Registrar decisões técnicas não óbvias (DuckDB vs cloud, escolha do modelo, grão "
     "do fato) como ADRs curtos em docs/.",
     "Principais decisões documentadas com contexto, alternativas e justificativa.",
     "S08-T01",
     "Revisão de conteúdo dos ADRs.",
     "Baixa"),
    (8, "S08-T04", "Documentação",
     "Screenshots/GIF e finalização do portfólio",
     "Adicionar GIF/screenshots do app ao README, sanear caminhos absolutos e dados "
     "pessoais, e preparar para publicação (alinhado à skill portfolio-docs).",
     "README com mídia; sem caminhos absolutos/dados pessoais; repositório pronto para "
     "ser público.",
     "S08-T02, S08-T03",
     "Revisão final: checklist de portfólio (mídia, reprodução, ausência de segredos).",
     "Média"),
]


# ============================================================
# Plano de Testes (resumo por feature)
# ============================================================

TESTES = [
    # (sprint, modulo, tipo_teste, descricao, sprint_alvo, status_inicial)
    (1, "core", "Smoke", "Import dos módulos do projeto", 1, "Pendente"),
    (1, "core", "Integração", "make setup/lint/test verdes em pasta limpa", 1, "Pendente"),

    (2, "ingestion", "Unitário", "Parsing GTFS com fixture reduzida + schema esperado", 2, "Pendente"),
    (2, "ingestion", "Unitário", "Parsing INMET (clima) e tipagem de data/numéricos", 2, "Pendente"),
    (2, "ingestion", "Unitário", "Parsing SIDRA (população) mockado", 2, "Pendente"),
    (2, "ingestion", "Integração", "Carga RAW no DuckDB e idempotência (rodar 2x)", 2, "Pendente"),
    (2, "ingestion", "Regressão", "Schema das fontes (falha clara se origem mudar)", 2, "Pendente"),

    (3, "dbt", "Integração", "dbt debug/parse verde", 3, "Pendente"),
    (3, "dbt", "Integração", "Testes not_null/unique nos staging", 3, "Pendente"),
    (3, "dbt", "Integração", "Relationships intermediate × staging (sem fan-out)", 3, "Pendente"),
    (3, "dbt", "Integração", "Unique/not_null nas dimensões + relationships do fato", 3, "Pendente"),
    (3, "dbt", "Regressão", "dbt build completo verde no CI", 3, "Pendente"),

    (4, "analysis", "Regressão", "Cálculo da métrica de atraso/demanda (casos manuais)", 4, "Pendente"),
    (4, "analysis", "Integração", "Queries de gargalos/sazonalidade com período vazio", 4, "Pendente"),
    (4, "analysis", "Unitário", "Agregado clima × atraso (sem chuva vs com chuva)", 4, "Pendente"),

    (5, "ml", "Unitário", "Shape/colunas do dataset de features", 5, "Pendente"),
    (5, "ml", "Regressão", "Ausência de vazamento temporal (corte de data)", 5, "Pendente"),
    (5, "ml", "Unitário", "Determinismo do dataset/treino com seed", 5, "Pendente"),
    (5, "ml", "Integração", "Baseline reporta métricas em fixture", 5, "Pendente"),
    (5, "ml", "Integração", "Treino registra run no MLflow (tracking temporário)", 5, "Pendente"),
    (5, "ml", "Integração", "Modelo serializado recarrega e prevê amostra estável", 5, "Pendente"),

    (6, "app", "Smoke", "Cada página do Streamlit importa sem erro", 6, "Pendente"),
    (6, "app", "Unitário", "Funções de agregação dos gráficos (entrada→saída)", 6, "Pendente"),
    (6, "app", "Unitário", "Preparação da camada do mapa", 6, "Pendente"),
    (6, "app", "Integração", "Função de predição do app com artefato de fixture", 6, "Pendente"),

    (7, "pipeline", "Integração", "Flow Prefect end-to-end em modo reduzido", 7, "Pendente"),
    (7, "ci", "Integração", "CI roda lint + pytest + dbt build (gates)", 7, "Pendente"),
    (7, "deploy", "Smoke", "docker build + run executa make all reduzido", 7, "Pendente"),
    (7, "global", "Cobertura", "pytest --cov >= 80% global", 7, "Pendente"),

    (8, "docs", "Manual", "Reproduzir projeto seguindo o README do zero", 8, "Pendente"),
    (8, "deploy", "Smoke", "App público carrega dashboard, mapa e previsão", 8, "Pendente"),
]


# ============================================================
# Construção da planilha
# ============================================================

def style_header(cell):
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = BORDER


def style_body(cell):
    cell.alignment = WRAP
    cell.border = BORDER


def set_columns(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


# ---- Aba Instruções ----

def build_instructions(wb):
    ws = wb.create_sheet("Instruções")
    set_columns(ws, [110])

    blocks = [
        ("Observatório de Mobilidade — Documentação Verdade das Sprints", "title"),
        ("Esta planilha é a fonte única de verdade para a execução do projeto. O chat do "
         "Claude Code que vai trabalhar nas tarefas deve usá-la como roteiro, registrar "
         "evolução nela e nunca tomar decisões de escopo fora dela.", "note"),
        ("", None),
        ("Como o Claude Code deve usar esta planilha", "section"),
        ("1. Sempre comece a sessão lendo a aba 'Visão Geral' para entender o estado macro.", "body"),
        ("2. Em seguida, abra a aba 'Backlog' e localize a próxima tarefa com status "
         "'Pendente' respeitando a coluna 'Dependências' (não pular ordem).", "body"),
        ("3. Marque a tarefa como 'Em andamento' e preencha 'Data Início'.", "body"),
        ("4. Implemente a tarefa seguindo exatamente o que está em 'Descrição' e "
         "'Critérios de Aceitação'. Não adicionar features extras.", "body"),
        ("5. Escrever os testes listados em 'Testes Obrigatórios' ANTES de marcar como "
         "concluída. Sem testes verdes a tarefa não é concluída.", "body"),
        ("6. Rodar a suíte completa de testes (pytest + dbt test quando aplicável) para "
         "garantir que nada regrediu antes de concluir.", "body"),
        ("7. Mudar status para 'Em teste' quando o código estiver pronto e os testes "
         "estiverem rodando; mudar para 'Concluída' quando tudo passar.", "body"),
        ("8. Preencher 'Data Fim' e qualquer 'Observação' relevante (decisões tomadas, "
         "arquivos criados, comandos úteis para rodar).", "body"),
        ("9. Se houver dúvida que mude o escopo ou comportamento esperado da tarefa, "
         "marcar 'Impedimento? = Sim', escrever a dúvida e PARAR. Perguntar para o "
         "humano via mensagem. Nunca decidir sozinho.", "body"),
        ("", None),
        ("Regras inegociáveis", "section"),
        ("• Toda tarefa que toca regra de negócio (cálculo da métrica de atraso/demanda, "
         "construção de features, idempotência de ingestão) DEVE ter testes que travam a "
         "regressão.", "body"),
        ("• A bateria de regressão da métrica de atraso/demanda (S04-T01) e do dataset de "
         "ML sem vazamento temporal (S05-T01) deve continuar verde em todas as sprints "
         "seguintes. Se quebrar, parar e investigar.", "body"),
        ("• Cobertura de testes do código Python deve permanecer ≥ 80% a partir da Sprint 7.", "body"),
        ("• Nunca commitar dados brutos pesados nem segredos. Sempre usar .env e .env.example "
         "e manter data/ no .gitignore.", "body"),
        ("• Transformações de dados sempre via dbt (camadas) — não criar marts manualmente "
         "fora do dbt.", "body"),
        ("• Ingestões devem ser idempotentes — reexecutar não pode duplicar dados.", "body"),
        ("• Para qualquer decisão técnica não trivial, registrar na aba 'Decisões'.", "body"),
        ("", None),
        ("Sobre a aba 'Plano de Testes'", "section"),
        ("Lista de checks de teste que precisam existir ao final do projeto. Marcar como "
         "'Concluída' assim que o teste correspondente estiver no repositório e verde no CI.", "body"),
        ("", None),
        ("Sobre a aba 'Impedimentos'", "section"),
        ("Registrar QUALQUER bloqueio antes de prosseguir. Cada linha tem ID, sprint/tarefa "
         "relacionada, descrição do problema, o que tentou e a pergunta para o humano.", "body"),
        ("", None),
        ("Sobre a aba 'Decisões'", "section"),
        ("Registrar decisões técnicas não óbvias: por que escolheu biblioteca X, por que "
         "estrutura Y, grão do fato, etc. Garante que o próximo Claude entenda o porquê.", "body"),
        ("", None),
        ("Convenções", "section"),
        ("• IDs de tarefa no formato S##-T## (ex.: S05-T03).", "body"),
        ("• Status válidos: Pendente, Em andamento, Bloqueada, Em teste, Concluída, "
         "Cancelada (dropdown na coluna Status).", "body"),
        ("• Datas no formato AAAA-MM-DD.", "body"),
        ("• Observações curtas e úteis, sem narrativa longa.", "body"),
    ]

    row = 1
    for text, kind in blocks:
        cell = ws.cell(row=row, column=1, value=text)
        if kind == "title":
            cell.font = TITLE_FONT
            ws.row_dimensions[row].height = 28
        elif kind == "section":
            cell.font = SECTION_FONT
            ws.row_dimensions[row].height = 22
        elif kind == "note":
            cell.font = NOTE_FONT
            cell.alignment = WRAP
            ws.row_dimensions[row].height = 45
        else:
            cell.alignment = WRAP
            ws.row_dimensions[row].height = 30
        row += 1


# ---- Aba Visão Geral ----

def build_overview(wb):
    ws = wb.create_sheet("Visão Geral")
    headers = ["Sprint", "Nome", "Objetivo", "Entregáveis", "Status", "Data Início", "Data Fim", "Observações"]
    widths = [8, 50, 60, 60, 16, 14, 14, 50]
    set_columns(ws, widths)

    for col, h in enumerate(headers, start=1):
        style_header(ws.cell(row=1, column=col, value=h))

    for i, s in enumerate(SPRINTS, start=2):
        ws.cell(row=i, column=1, value=s["numero"])
        ws.cell(row=i, column=2, value=s["nome"])
        ws.cell(row=i, column=3, value=s["objetivo"])
        ws.cell(row=i, column=4, value=s["entregaveis"])
        ws.cell(row=i, column=5, value="Pendente")
        ws.cell(row=i, column=6, value="")
        ws.cell(row=i, column=7, value="")
        ws.cell(row=i, column=8, value="")
        for c in range(1, 9):
            cell = ws.cell(row=i, column=c)
            style_body(cell)
            if c == 1:
                cell.fill = SPRINT_FILL

        ws.row_dimensions[i].height = 75

    last_row = 1 + len(SPRINTS)
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True)
    dv.add(f"E2:E{last_row}")
    ws.add_data_validation(dv)

    ws.freeze_panes = "A2"


# ---- Aba Backlog ----

def build_backlog(wb):
    ws = wb.create_sheet("Backlog")
    headers = [
        "Sprint", "ID", "Tipo", "Prioridade", "Tarefa", "Descrição",
        "Critérios de Aceitação", "Dependências", "Testes Obrigatórios",
        "Status", "Data Início", "Data Fim", "Observações",
        "Impedimento?", "Detalhe do Impedimento",
    ]
    widths = [8, 10, 14, 11, 38, 60, 55, 18, 55, 14, 13, 13, 40, 13, 40]
    set_columns(ws, widths)

    for col, h in enumerate(headers, start=1):
        style_header(ws.cell(row=1, column=col, value=h))

    for i, t in enumerate(TAREFAS, start=2):
        sprint, tid, tipo, titulo, desc, criterios, deps, testes, prio = t
        ws.cell(row=i, column=1, value=sprint)
        ws.cell(row=i, column=2, value=tid)
        ws.cell(row=i, column=3, value=tipo)
        ws.cell(row=i, column=4, value=prio)
        ws.cell(row=i, column=5, value=titulo)
        ws.cell(row=i, column=6, value=desc)
        ws.cell(row=i, column=7, value=criterios)
        ws.cell(row=i, column=8, value=deps)
        ws.cell(row=i, column=9, value=testes)
        ws.cell(row=i, column=10, value="Pendente")
        ws.cell(row=i, column=11, value="")
        ws.cell(row=i, column=12, value="")
        ws.cell(row=i, column=13, value="")
        ws.cell(row=i, column=14, value="Não")
        ws.cell(row=i, column=15, value="")

        for c in range(1, 16):
            style_body(ws.cell(row=i, column=c))

        ws.row_dimensions[i].height = 110

    last_row = 1 + len(TAREFAS)

    dv_status = DataValidation(type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True)
    dv_status.add(f"J2:J{last_row}")
    ws.add_data_validation(dv_status)

    dv_prio = DataValidation(type="list", formula1=f'"{",".join(PRIORIDADE_OPTIONS)}"', allow_blank=True)
    dv_prio.add(f"D2:D{last_row}")
    ws.add_data_validation(dv_prio)

    dv_tipo = DataValidation(type="list", formula1=f'"{",".join(TIPO_OPTIONS)}"', allow_blank=True)
    dv_tipo.add(f"C2:C{last_row}")
    ws.add_data_validation(dv_tipo)

    dv_imp = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    dv_imp.add(f"N2:N{last_row}")
    ws.add_data_validation(dv_imp)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:O{last_row}"


# ---- Aba Plano de Testes ----

def build_tests(wb):
    ws = wb.create_sheet("Plano de Testes")
    headers = ["Sprint Alvo", "Módulo", "Tipo de Teste", "Descrição", "Arquivo de Teste", "Status", "Última Execução", "Observações"]
    widths = [12, 16, 14, 70, 40, 14, 18, 40]
    set_columns(ws, widths)

    for col, h in enumerate(headers, start=1):
        style_header(ws.cell(row=1, column=col, value=h))

    for i, t in enumerate(TESTES, start=2):
        sprint, modulo, tipo, desc, sprint_alvo, status = t
        ws.cell(row=i, column=1, value=sprint_alvo)
        ws.cell(row=i, column=2, value=modulo)
        ws.cell(row=i, column=3, value=tipo)
        ws.cell(row=i, column=4, value=desc)
        ws.cell(row=i, column=5, value="")
        ws.cell(row=i, column=6, value=status)
        ws.cell(row=i, column=7, value="")
        ws.cell(row=i, column=8, value="")
        for c in range(1, 9):
            style_body(ws.cell(row=i, column=c))
        ws.row_dimensions[i].height = 35

    last_row = 1 + len(TESTES)
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True)
    dv.add(f"F2:F{last_row}")
    ws.add_data_validation(dv)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{last_row}"


# ---- Aba Impedimentos ----

def build_blockers(wb):
    ws = wb.create_sheet("Impedimentos")
    headers = ["ID", "Data", "Sprint", "Tarefa Relacionada", "Descrição", "O que tentou", "Pergunta para o humano", "Status", "Resposta / Decisão"]
    widths = [8, 14, 8, 14, 50, 50, 50, 16, 50]
    set_columns(ws, widths)

    for col, h in enumerate(headers, start=1):
        style_header(ws.cell(row=1, column=col, value=h))

    for i in range(2, 22):
        for c in range(1, 10):
            style_body(ws.cell(row=i, column=c))
        ws.row_dimensions[i].height = 35

    dv = DataValidation(type="list", formula1='"Aberto,Aguardando humano,Resolvido,Descartado"', allow_blank=True)
    dv.add("H2:H200")
    ws.add_data_validation(dv)

    ws.freeze_panes = "A2"


# ---- Aba Decisões ----

def build_decisions(wb):
    ws = wb.create_sheet("Decisões")
    headers = ["ID", "Data", "Sprint", "Contexto", "Decisão", "Alternativas consideradas", "Justificativa", "Impacto / Notas"]
    widths = [8, 14, 8, 50, 50, 50, 50, 40]
    set_columns(ws, widths)

    for col, h in enumerate(headers, start=1):
        style_header(ws.cell(row=1, column=col, value=h))

    for i in range(2, 22):
        for c in range(1, 9):
            style_body(ws.cell(row=i, column=c))
        ws.row_dimensions[i].height = 35

    ws.freeze_panes = "A2"


# ============================================================
# Build
# ============================================================

def main():
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    build_instructions(wb)
    build_overview(wb)
    build_backlog(wb)
    build_tests(wb)
    build_blockers(wb)
    build_decisions(wb)

    out_path = "observatorio_mobilidade_planejamento_sprints.xlsx"
    wb.save(out_path)
    print(f"Planilha gerada: {out_path}")
    print(f"Total de sprints: {len(SPRINTS)}")
    print(f"Total de tarefas: {len(TAREFAS)}")
    print(f"Total de itens de teste: {len(TESTES)}")


if __name__ == "__main__":
    main()
