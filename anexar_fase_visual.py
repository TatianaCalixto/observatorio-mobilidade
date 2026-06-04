"""Anexa a próxima fase (Redesign Visual & UX) à planilha de sprints EXISTENTE.

Importante: NÃO regenera a planilha do zero — abre a planilha já preenchida pelo
executor (status, decisões, impedimentos) e apenas ACRESCENTA as novas sprints,
tarefas e itens de teste, preservando tudo que já existe.

As novas tarefas vêm do documento docs/sugestoes_melhorias_visuais_app.md.
"""

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ARQUIVO = r"d:\Projetos\Dados\observatorio_mobilidade_planejamento_sprints.xlsx"

# ---- estilos (iguais aos do gerador original) ----
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SPRINT_FILL = PatternFill("solid", fgColor="D9E1F2")

STATUS_OPTIONS = ["Pendente", "Em andamento", "Bloqueada", "Em teste", "Concluída", "Cancelada"]
PRIORIDADE_OPTIONS = ["Alta", "Média", "Baixa"]
TIPO_OPTIONS = ["Backend", "Frontend", "Mobile", "Infra", "Banco", "Teste", "Documentação", "Deploy"]


# ============================================================
# Conteúdo da nova fase (derivado do .md de sugestões)
# ============================================================

SPRINTS_NOVAS = [
    {
        "numero": 9,
        "nome": "Sprint 9 — Identidade Visual & Gráficos",
        "objetivo": "Dar identidade ao app (tema de cores/fonte) e redesenhar os gráficos "
                    "com uma biblioteca que permita controlar cor, eixos, unidades e "
                    "tooltips. Resolve a maior parte da percepção de 'visual cru'.",
        "entregaveis": "Tema aplicado (.streamlit/config.toml), módulo de gráficos "
                       "reutilizável com paleta semântica, gráficos de linha e barra "
                       "redesenhados, sem alterar a lógica de dados.",
    },
    {
        "numero": 10,
        "nome": "Sprint 10 — Mapa, Tabelas & Polimento de UX",
        "objetivo": "Sofisticar o mapa (escala de cor moderna, basemap, legenda), melhorar "
                    "tabelas e KPIs e organizar a navegação (abas, divisórias), incluindo "
                    "a página de Previsões.",
        "entregaveis": "Mapa com escala perceptual + legenda, tabela de gargalos com "
                       "barras inline, KPIs com variação, página de KPIs em abas e página "
                       "de Previsões mais ilustrativa.",
    },
]

# (sprint, id, tipo, titulo, descricao, criterios, dependencias, testes, prioridade)
TAREFAS_NOVAS = [
    # ---------------- Sprint 9 ----------------
    (9, "S09-T01", "Frontend",
     "Tema visual do app (cores e fonte)",
     "Criar .streamlit/config.toml definindo primaryColor (sugestão: azul/petróleo de "
     "transporte), backgroundColor, secondaryBackgroundColor, textColor e font. Tirar o "
     "visual padrão do Streamlit.",
     "App carrega com a nova paleta e fonte; nenhuma mudança na lógica de dados; "
     "aparência consistente em todas as páginas.",
     "S06-T01",
     "Smoke: app sobe com o config sem erro; revisão visual das 4 páginas.",
     "Alta"),
    (9, "S09-T02", "Frontend",
     "Módulo de gráficos + paleta semântica (decisão técnica)",
     "Escolher a biblioteca de gráficos (Altair ou Plotly) e registrar a escolha como "
     "ADR/Decisão. Criar app/charts.py com paleta semântica fixa (ex.: chuva=azul, "
     "sem chuva=âmbar; dia útil vs fim de semana) e helpers de eixo/tooltip reutilizáveis.",
     "Módulo de gráficos central criado; paleta documentada e reutilizável; decisão "
     "registrada na aba Decisões.",
     "S09-T01",
     "Unit tests dos helpers (entrada→spec do gráfico) com dados de fixture.",
     "Alta"),
    (9, "S09-T03", "Frontend",
     "Redesenhar gráficos de linha (oferta ao longo do tempo)",
     "Substituir st.line_chart (Visão Geral e KPIs) por gráficos do novo módulo: eixos "
     "com rótulo e unidade ('viagens/dia'), datas legíveis, tooltip formatado, área "
     "preenchida e cor da marca.",
     "Gráficos de linha exibem título de eixo, unidade e tooltip formatado; usam a "
     "paleta; estado vazio preservado.",
     "S09-T02",
     "Unit test da função que monta o gráfico de linha (colunas/encodings esperados).",
     "Alta"),
    (9, "S09-T04", "Frontend",
     "Redesenhar gráficos de barra (sazonalidade e chuva)",
     "Substituir st.bar_chart por barras do novo módulo: dias da semana ordenados "
     "Seg→Dom, cor distinta para dia útil vs fim de semana; gráfico de chuva com cores "
     "semânticas (com/sem chuva) e rótulos de valor.",
     "Barras ordenadas corretamente, com cores semânticas e rótulos; tooltips claros.",
     "S09-T02",
     "Unit test da preparação dos dados das barras (ordem Seg→Dom; mapeamento de cor).",
     "Média"),
    (9, "S09-T05", "Teste",
     "Cobertura dos novos gráficos e smoke do app",
     "Garantir testes para as funções de app/charts.py e smoke de que cada página ainda "
     "importa e renderiza sem erro. Manter cobertura ≥ 80%.",
     "Funções de gráfico cobertas por teste; smoke das 4 páginas verde; cobertura ≥ 80%.",
     "S09-T03, S09-T04",
     "Smoke por página + unit tests de charts; relatório de cobertura no CI.",
     "Alta"),

    # ---------------- Sprint 10 ----------------
    (10, "S10-T01", "Frontend",
     "Mapa com escala de cor moderna + basemap + legenda",
     "Trocar o gradiente vermelho→amarelo por uma escala perceptual (ex.: Viridis), "
     "colorblind-safe; adicionar basemap claro/escuro (map_style) e uma legenda "
     "explicando o que cor e tamanho dos pontos representam.",
     "Mapa usa escala perceptual; basemap aplicado; legenda visível e correta.",
     "S06-T03",
     "Unit test da função que mapeia intensidade→cor (faixas esperadas).",
     "Média"),
    (10, "S10-T02", "Frontend",
     "Opção de mapa de calor para densidade",
     "Adicionar alternância (toggle) entre pontos e HeatmapLayer/HexagonLayer para quando "
     "há muitas paradas sobrepostas.",
     "Toggle funciona; camada de calor renderiza sem erro; pontos continuam disponíveis.",
     "S10-T01",
     "Unit test da preparação dos dados da camada de calor.",
     "Baixa"),
    (10, "S10-T03", "Frontend",
     "Tabela de gargalos com barras e formatação",
     "Usar st.dataframe(column_config=...) com ProgressColumn para o headway e NumberColumn "
     "com formatação de número nas demais colunas.",
     "Tabela mostra barra inline para headway e números formatados; ordenação preservada.",
     "S09-T01",
     "Unit test da preparação do DataFrame da tabela (colunas/formatos).",
     "Média"),
    (10, "S10-T04", "Frontend",
     "KPIs com variação e cartões",
     "Adicionar delta nos st.metric (variação vs período anterior) e agrupar em "
     "st.container(border=True) para efeito de cartão.",
     "KPIs mostram seta de variação coerente; layout em cartões; cálculo do delta testado.",
     "S09-T01",
     "Unit test do cálculo de variação (período atual vs anterior).",
     "Média"),
    (10, "S10-T05", "Frontend",
     "Reorganizar página de KPIs em abas",
     "Quebrar a página 'KPIs & Análises' em abas (st.tabs): Tendência, Sazonalidade, "
     "Gargalos; adicionar st.divider() entre seções para respiro visual.",
     "Página em abas reduz rolagem; conteúdo equivalente distribuído; sem regressão de "
     "filtros.",
     "S09-T03, S09-T04",
     "Smoke da página em abas (cada aba importa e renderiza).",
     "Baixa"),
    (10, "S10-T06", "Frontend",
     "Página de Previsões mais ilustrativa",
     "Complementar o número único com um gauge (velocímetro) ou mini-gráfico de "
     "sensibilidade (demanda × chuva). Manter o aviso de 'demanda simulada' (DEC-007).",
     "Previsão exibida com visual ilustrativo; aviso de simulação mantido; sem mudança "
     "no modelo.",
     "S09-T02",
     "Unit test da função que gera os pontos de sensibilidade do gráfico.",
     "Baixa"),
    (10, "S10-T07", "Teste",
     "Polimento final, smoke e cobertura",
     "Smoke completo das 4 páginas após o redesign, revisão de consistência visual e "
     "manutenção da cobertura ≥ 80%.",
     "Todas as páginas renderizam; visual consistente; cobertura ≥ 80%; baterias de "
     "regressão de dados/ML seguem verdes.",
     "S10-T01, S10-T03, S10-T04, S10-T05, S10-T06",
     "Smoke das 4 páginas + reexecução das regressões críticas; cobertura no CI.",
     "Alta"),
]

# (modulo, tipo_teste, descricao, sprint_alvo, status)
TESTES_NOVOS = [
    ("app/charts", "Unitário", "Helpers de gráfico geram encodings/cores esperados", 9, "Pendente"),
    ("app/charts", "Unitário", "Gráfico de linha: rótulos de eixo e unidade corretos", 9, "Pendente"),
    ("app/charts", "Unitário", "Barras de sazonalidade ordenadas Seg→Dom + cores", 9, "Pendente"),
    ("app", "Smoke", "4 páginas importam/renderizam com o novo tema", 9, "Pendente"),
    ("app/mapa", "Unitário", "Mapeamento intensidade→cor (escala perceptual)", 10, "Pendente"),
    ("app/mapa", "Unitário", "Preparação da camada de mapa de calor", 10, "Pendente"),
    ("app", "Unitário", "Tabela de gargalos: colunas e formatos do column_config", 10, "Pendente"),
    ("app", "Unitário", "Cálculo de variação (delta) dos KPIs", 10, "Pendente"),
    ("app", "Smoke", "Página de KPIs em abas renderiza cada aba", 10, "Pendente"),
    ("app", "Unitário", "Pontos de sensibilidade da página de Previsões", 10, "Pendente"),
    ("global", "Cobertura", "pytest --cov >= 80% após o redesign", 10, "Pendente"),
]


def style_body(cell):
    cell.alignment = WRAP
    cell.border = BORDER


def main():
    wb = load_workbook(ARQUIVO)

    # nomes de aba com acento (resolvidos pelo objeto, não por string fixa)
    ws_vg = next(ws for ws in wb.worksheets if ws.title.startswith("Vis"))
    ws_bk = wb["Backlog"]
    ws_pt = wb["Plano de Testes"]

    # ---- Visão Geral: anexar sprints ----
    start_vg = ws_vg.max_row + 1
    for off, s in enumerate(SPRINTS_NOVAS):
        r = start_vg + off
        ws_vg.cell(r, 1, s["numero"])
        ws_vg.cell(r, 2, s["nome"])
        ws_vg.cell(r, 3, s["objetivo"])
        ws_vg.cell(r, 4, s["entregaveis"])
        ws_vg.cell(r, 5, "Pendente")
        for c in range(1, 9):
            style_body(ws_vg.cell(r, c))
            if c == 1:
                ws_vg.cell(r, c).fill = SPRINT_FILL
        ws_vg.row_dimensions[r].height = 75
    end_vg = start_vg + len(SPRINTS_NOVAS) - 1
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True)
    dv.add(f"E{start_vg}:E{end_vg}")
    ws_vg.add_data_validation(dv)

    # ---- Backlog: anexar tarefas ----
    start_bk = ws_bk.max_row + 1
    for off, t in enumerate(TAREFAS_NOVAS):
        r = start_bk + off
        sprint, tid, tipo, titulo, desc, criterios, deps, testes, prio = t
        ws_bk.cell(r, 1, sprint)
        ws_bk.cell(r, 2, tid)
        ws_bk.cell(r, 3, tipo)
        ws_bk.cell(r, 4, prio)
        ws_bk.cell(r, 5, titulo)
        ws_bk.cell(r, 6, desc)
        ws_bk.cell(r, 7, criterios)
        ws_bk.cell(r, 8, deps)
        ws_bk.cell(r, 9, testes)
        ws_bk.cell(r, 10, "Pendente")
        ws_bk.cell(r, 14, "Não")
        for c in range(1, 16):
            style_body(ws_bk.cell(r, c))
        ws_bk.row_dimensions[r].height = 110
    end_bk = start_bk + len(TAREFAS_NOVAS) - 1
    for col, opts in [("J", STATUS_OPTIONS), ("D", PRIORIDADE_OPTIONS), ("C", TIPO_OPTIONS)]:
        d = DataValidation(type="list", formula1=f'"{",".join(opts)}"', allow_blank=True)
        d.add(f"{col}{start_bk}:{col}{end_bk}")
        ws_bk.add_data_validation(d)
    d_imp = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    d_imp.add(f"N{start_bk}:N{end_bk}")
    ws_bk.add_data_validation(d_imp)
    ws_bk.auto_filter.ref = f"A1:O{end_bk}"

    # ---- Plano de Testes: anexar itens ----
    start_pt = ws_pt.max_row + 1
    for off, t in enumerate(TESTES_NOVOS):
        r = start_pt + off
        modulo, tipo, desc, sprint_alvo, status = t
        ws_pt.cell(r, 1, sprint_alvo)
        ws_pt.cell(r, 2, modulo)
        ws_pt.cell(r, 3, tipo)
        ws_pt.cell(r, 4, desc)
        ws_pt.cell(r, 6, status)
        for c in range(1, 9):
            style_body(ws_pt.cell(r, c))
        ws_pt.row_dimensions[r].height = 35
    end_pt = start_pt + len(TESTES_NOVOS) - 1
    d = DataValidation(type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True)
    d.add(f"F{start_pt}:F{end_pt}")
    ws_pt.add_data_validation(d)
    ws_pt.auto_filter.ref = f"A1:H{end_pt}"

    wb.save(ARQUIVO)
    print("Planilha atualizada (fase visual anexada).")
    print(f"  Sprints novas: {len(SPRINTS_NOVAS)} (linhas {start_vg}-{end_vg})")
    print(f"  Tarefas novas: {len(TAREFAS_NOVAS)} (linhas {start_bk}-{end_bk})")
    print(f"  Itens de teste novos: {len(TESTES_NOVOS)} (linhas {start_pt}-{end_pt})")


if __name__ == "__main__":
    main()
